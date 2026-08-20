"""Adapter for California OES hazardous-material spill notifications."""

from __future__ import annotations

import csv
import heapq
import re
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Iterator

from .models import UNITED_STATES_COUNTRY, IncidentRecord
from .taxonomy import classify_hazard, occupational_potential_severity
from .text import (
    clean_text,
    is_explicit_non_release,
    is_probably_english,
    sampling_rank,
    text_digest,
    word_count,
)


CALOES_SOURCE_URL = (
    "https://www.caloes.ca.gov/office-of-the-director/operations/"
    "response-operations/fire-rescue/hazardous-materials/spill-release-reporting/"
)
CALOES_SOURCE_LICENSE = (
    "Publicly downloadable California spill-notification data; no explicit "
    "dataset-wide licence was identified; preserve attribution and review "
    "caller-authored narrative terms before redistribution"
)

_HEADER_PATTERN = re.compile(r"[^a-z0-9]+")


def _header(value: object) -> str:
    return _HEADER_PATTERN.sub("", clean_text(value).casefold())


def _as_row(headers: list[str], values: list[object] | tuple[object, ...]) -> dict[str, object]:
    return {
        header: values[index] if index < len(values) else ""
        for index, header in enumerate(headers)
        if header
    }


def _iter_csv(path: Path) -> Iterator[dict[str, object]]:
    with path.open(encoding="utf-8-sig", errors="replace", newline="") as stream:
        reader = csv.reader(stream)
        try:
            headers = [_header(value) for value in next(reader)]
        except StopIteration:
            return
        for values in reader:
            yield _as_row(headers, values)


def _iter_xlsx(path: Path) -> Iterator[dict[str, object]]:
    try:
        import openpyxl
    except ImportError as error:  # pragma: no cover - exercised in minimal installs
        raise RuntimeError(
            "Reading Cal OES .xlsx files requires openpyxl>=3.1"
        ) from error

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        try:
            headers = [_header(value) for value in next(rows)]
        except StopIteration:
            return
        for values in rows:
            yield _as_row(headers, values)
    finally:
        workbook.close()


def _iter_xls(path: Path) -> Iterator[dict[str, object]]:
    try:
        import xlrd
    except ImportError as error:  # pragma: no cover - exercised in minimal installs
        raise RuntimeError("Reading Cal OES .xls files requires xlrd>=2.0") from error

    workbook = xlrd.open_workbook(path, on_demand=True)
    try:
        worksheet = workbook.sheet_by_index(0)
        if worksheet.nrows == 0:
            return
        headers = [_header(value) for value in worksheet.row_values(0)]
        for row_index in range(1, worksheet.nrows):
            yield _as_row(headers, worksheet.row_values(row_index))
    finally:
        workbook.release_resources()


def _iter_rows(path: Path) -> Iterator[dict[str, object]]:
    suffix = path.suffix.casefold()
    if suffix == ".csv":
        yield from _iter_csv(path)
    elif suffix == ".xlsx":
        yield from _iter_xlsx(path)
    elif suffix == ".xls":
        yield from _iter_xls(path)
    else:
        raise ValueError(f"Unsupported Cal OES source format: {path}")


def _discover_files(path: Path) -> list[Path]:
    if path.is_file():
        return [path]
    supported = {".csv", ".xls", ".xlsx"}
    return sorted(
        candidate
        for candidate in path.rglob("*")
        if candidate.is_file() and candidate.suffix.casefold() in supported
    )


def _value(row: dict[str, object], *names: str) -> object:
    for name in names:
        value = row.get(_header(name))
        if clean_text(value):
            return value
    return ""


def _source_id(value: object) -> str:
    return clean_text(value).lstrip("'").strip()


def _event_date(value: object) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    cleaned = clean_text(value)
    try:
        serial = float(cleaned)
    except ValueError:
        serial = 0.0
    if 20_000 <= serial <= 80_000:
        return (datetime(1899, 12, 30) + timedelta(days=serial)).date().isoformat()
    for date_format in (
        "%Y-%m-%d",
        "%m/%d/%Y",
        "%Y-%m-%d %H:%M:%S",
        "%m/%d/%Y %H:%M:%S",
    ):
        try:
            return datetime.strptime(cleaned, date_format).date().isoformat()
        except ValueError:
            pass
    return cleaned


def _best_event_date(row: dict[str, object]) -> str:
    primary = _event_date(_value(row, "Incident Date"))
    fallback = _event_date(_value(row, "Notified Date"))
    if _plausible_iso_date(primary):
        if _plausible_iso_date(fallback):
            incident = datetime.strptime(primary, "%Y-%m-%d").date()
            notified = datetime.strptime(fallback, "%Y-%m-%d").date()
            # A notification cannot precede the incident it reports.  Cal OES
            # contains a small number of obvious year-entry errors; use the
            # source notification date while retaining both raw values below.
            if incident > notified:
                return fallback
        return primary
    if _plausible_iso_date(fallback):
        return fallback
    return primary or fallback or "Date not reported"


def _plausible_iso_date(value: str) -> bool:
    try:
        parsed = datetime.strptime(value, "%Y-%m-%d").date()
    except ValueError:
        return False
    return date(1990, 1, 1) <= parsed <= date.today()


def _number(value: object) -> float:
    cleaned = clean_text(value).replace(",", "")
    try:
        return float(cleaned) if cleaned else 0.0
    except ValueError:
        return 0.0


def _truthy(value: object) -> bool:
    return clean_text(value).casefold() in {"yes", "y", "true", "1"}


def _actual_severity(row: dict[str, object]) -> str:
    fatalities_flag = _value(row, "Fatals YES/NO", "Fatalities YES/NO")
    fatalities_count = _value(row, "Fatals #", "Fatalities #")
    if _truthy(fatalities_flag) or _number(fatalities_count) > 0:
        return "Severe"

    injuries_flag = _value(row, "Injuries YES/NO")
    injuries_count = _value(row, "Injuries #")
    evacuations_flag = _value(row, "Evacs YES/NO", "Evacuations YES/NO")
    evacuations_count = _value(row, "Evacs #", "Evacuations #")
    if (
        _truthy(injuries_flag)
        or _number(injuries_count) > 0
        or _truthy(evacuations_flag)
        or _number(evacuations_count) > 0
    ):
        return "Medium"

    flags = [
        clean_text(value).casefold()
        for value in (fatalities_flag, injuries_flag, evacuations_flag)
        if clean_text(value)
    ]
    if flags and all(value in {"no", "n", "false", "0"} for value in flags):
        return "Low"
    return "Unknown"


def _severity_evidence(row: dict[str, object]) -> str:
    fields = (
        "Injuries YES/NO",
        "Injuries #",
        "Fatals YES/NO",
        "Fatals #",
        "Evacs YES/NO",
        "Evacs #",
    )
    return " | ".join(
        f"{field}={clean_text(_value(row, field))}"
        for field in fields
        if clean_text(_value(row, field))
    )


def _raw_hazard_evidence(
    row: dict[str, object],
    substances: list[str],
    substance_types: list[str],
    cause: str,
) -> str:
    fields = [
        f"substances={'; '.join(substances)}" if substances else "",
        f"substance_types={'; '.join(substance_types)}" if substance_types else "",
        "substance_quantities="
        + "; ".join(
            (
                f"{clean_text(_value(row, f'{index}. Quantity'))} "
                f"{clean_text(_value(row, f'{index}. Measure'))}"
            ).strip()
            for index in (1, 2, 3)
            if clean_text(_value(row, f"{index}. Quantity"))
        )
        if any(
            clean_text(_value(row, f"{index}. Quantity"))
            for index in (1, 2, 3)
        )
        else "",
        f"cause={cause}" if cause else "",
        f"contained={clean_text(_value(row, 'Contained'))}"
        if clean_text(_value(row, "Contained"))
        else "",
        f"water={clean_text(_value(row, 'Water?'))}"
        if clean_text(_value(row, "Water?"))
        else "",
        f"waterway={clean_text(_value(row, 'Water Way'))}"
        if clean_text(_value(row, "Water Way"))
        else "",
        f"known_impact={clean_text(_value(row, 'Known Impact'))}"
        if clean_text(_value(row, "Known Impact"))
        else "",
        f"spill_site={clean_text(_value(row, 'Spill Site'))}"
        if clean_text(_value(row, "Spill Site"))
        else "",
        f"location={clean_text(_value(row, 'Location'))}"
        if clean_text(_value(row, "Location"))
        else "",
        f"city={clean_text(_value(row, 'City'))}"
        if clean_text(_value(row, "City"))
        else "",
        f"county={clean_text(_value(row, 'County'))}"
        if clean_text(_value(row, "County"))
        else "",
        f"cleanup={clean_text(_value(row, 'Cleanup'))}"
        if clean_text(_value(row, "Cleanup"))
        else "",
        f"incident_date={clean_text(_value(row, 'Incident Date'))}"
        if clean_text(_value(row, "Incident Date"))
        else "",
        f"notified_date={clean_text(_value(row, 'Notified Date'))}"
        if clean_text(_value(row, "Notified Date"))
        else "",
    ]
    return " | ".join(value for value in fields if value)


def _push_bounded(
    heap: list[tuple[int, str, IncidentRecord]],
    capacity: int,
    record: IncidentRecord,
) -> None:
    if capacity <= 0:
        return
    item = (-record.sampling_rank, record.case_no, record)
    if len(heap) < capacity:
        heapq.heappush(heap, item)
    elif record.sampling_rank < -heap[0][0]:
        heapq.heapreplace(heap, item)


def build_caloes_environment_candidates(
    path: str,
    capacity: int,
    seed: int,
    minimum_words: int,
    maximum_rows: int | None = None,
) -> tuple[list[IncidentRecord], dict[str, int]]:
    """Build deterministic Environment candidates from Cal OES spill files."""
    source = "CALOES_SPILL"
    heap: list[tuple[int, str, IncidentRecord]] = []
    seen_source_ids: set[str] = set()
    statistics = {
        "source_files_read": 0,
        "rows_read": 0,
        "missing_source_id": 0,
        "empty_or_short": 0,
        "non_english": 0,
        "duplicate_source_ids": 0,
        "explicit_non_release": 0,
        "non_environmental_rail_event": 0,
        "eligible_unique": 0,
    }
    stop = False

    for source_path in _discover_files(Path(path)):
        if stop:
            break
        statistics["source_files_read"] += 1
        for row in _iter_rows(source_path):
            if maximum_rows is not None and statistics["rows_read"] >= maximum_rows:
                stop = True
                break
            statistics["rows_read"] += 1

            source_id = _source_id(_value(row, "Control#", "Control Number"))
            if not source_id:
                statistics["missing_source_id"] += 1
                continue
            if source_id in seen_source_ids:
                statistics["duplicate_source_ids"] += 1
                continue

            description = clean_text(_value(row, "Description"))
            if word_count(description) < minimum_words:
                statistics["empty_or_short"] += 1
                continue
            if not is_probably_english(description, minimum_source_words=5):
                statistics["non_english"] += 1
                continue

            substances = [
                clean_text(_value(row, f"{index}. Substance"))
                for index in (1, 2, 3)
                if clean_text(_value(row, f"{index}. Substance"))
            ]
            substance_types = [
                clean_text(_value(row, f"{index}. Type"))
                for index in (1, 2, 3)
                if clean_text(_value(row, f"{index}. Type"))
            ]
            positive_quantity = any(
                _number(_value(row, f"{index}. Quantity")) > 0
                for index in (1, 2, 3)
            )
            water = clean_text(_value(row, "Water?"))
            waterway = clean_text(_value(row, "Water Way"))
            known_impact = clean_text(_value(row, "Known Impact"))
            known_impact_present = known_impact.casefold() not in {
                "",
                "n",
                "no",
                "none",
                "not applicable",
                "n/a",
                "unknown",
            }
            positive_environment_evidence = (
                positive_quantity
                or _truthy(water)
                or waterway.casefold() not in {"", "n", "no", "none", "n/a"}
            )
            if is_explicit_non_release(description) and not positive_environment_evidence:
                statistics["explicit_non_release"] += 1
                continue
            railroad_only = bool(substance_types) and all(
                "railroad" in value.casefold() for value in substance_types
            )
            non_material_rail_subject = bool(substances) and all(
                any(
                    term in value.casefold()
                    for term in ("train vs", "trespass", "derailment")
                )
                for value in substances
            )
            if railroad_only and non_material_rail_subject:
                statistics["non_environmental_rail_event"] += 1
                continue
            seen_source_ids.add(source_id)

            cause = clean_text(_value(row, "CAUSE"))
            cause_other = clean_text(_value(row, "CAUSE_Other"))
            cause_text = " — ".join(value for value in (cause, cause_other) if value)
            if (
                _truthy(water)
                or waterway.casefold() not in {"", "n", "no", "none"}
                or known_impact_present
            ):
                hazard_type = "Environmental Pollution"
            else:
                hazard_type = classify_hazard(
                    "spill release",
                    *substances,
                    *substance_types,
                    cause_text,
                    description,
                )
                if hazard_type == "Other/Unknown":
                    hazard_type = "Loss of Containment/Release"

            subject = ", ".join(substances[:2]) or "Hazardous-material"
            title = f"{subject} spill notification"
            if cause_text and cause_text.casefold() != "unknown":
                title += f" — {cause_text.title()}"
            hazard = " — ".join(
                value for value in (", ".join(substances[:3]), cause_text) if value
            ) or "Reported hazardous-material spill"
            actual_severity = _actual_severity(row)
            record = IncidentRecord(
                case_no=f"CALOES:{source_id}",
                country=UNITED_STATES_COUNTRY,
                title=title,
                description=description,
                case_type="Environment",
                hazard=hazard,
                hazard_type=hazard_type,
                actual_severity=actual_severity,
                potential_severity=occupational_potential_severity(
                    actual_severity,
                    hazard_type,
                ),
                source=source,
                source_record_id=source_id,
                source_url=CALOES_SOURCE_URL,
                source_license=CALOES_SOURCE_LICENSE,
                event_date=_best_event_date(row),
                raw_case_type=(
                    "Cal OES spill notification"
                    + (
                        f"; spill_site={clean_text(_value(row, 'Spill Site'))}"
                        if clean_text(_value(row, "Spill Site"))
                        else ""
                    )
                ),
                raw_hazard=_raw_hazard_evidence(
                    row,
                    substances,
                    substance_types,
                    cause_text,
                ),
                raw_severity=(
                    _severity_evidence(row)
                    or "No source consequence fields reported"
                ),
                title_provenance="deterministic_from_source_substance_and_cause",
                hazard_label_method="rule_from_source_spill_water_substance_and_cause",
                actual_severity_label_method="source_consequence_fields_rule",
                potential_severity_label_method="documented_hazard_matrix_rule",
                text_hash=text_digest(description),
                sampling_rank=sampling_rank(seed, source, source_id),
            )
            statistics["eligible_unique"] += 1
            _push_bounded(heap, capacity, record)

    records = sorted(
        (item[2] for item in heap),
        key=lambda record: (record.sampling_rank, record.case_no),
    )
    statistics["candidates_returned"] = len(records)
    return records, statistics
